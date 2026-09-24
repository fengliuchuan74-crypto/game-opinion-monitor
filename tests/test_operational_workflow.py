import copy
import json
import tempfile
import unittest
from datetime import datetime,timezone,timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock,patch

import pandas as pd
from openpyxl import load_workbook

from collectors.app_store import AppStoreCollector
from collectors.base import RawReview,CollectionResult
from modules.data_loader import prepare_dataframe
from modules.deliverables import excel_bytes,build_report,workbook
from modules.imports import parse_import
from modules.issue_classifier import classify_issue
from modules.monitoring import configure_target,request_collection,claim_job,perform_job,target_state,recent_runs
from modules.operations import analyze_reviews,snapshot,issue_plans,review_override,create_action,actions_for,update_action,LOCAL_TZ
from modules.review_exports import export_reviews_by_game
from modules.review_store import init_db,save_reviews,read_reviews
from modules.sentiment import analyze_sentiment
from modules.storage import database,backup_database


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder=Path(self.temporary.name)
        self.db=self.folder/'reviews.sqlite3'
        init_db(self.db)
        self.now=datetime(2026,9,22,12,tzinfo=LOCAL_TZ)

    def review(self,identity='1',**changes):
        values=dict(platform='App Store',external_id=identity,date='2026-09-22T01:00:00+00:00',
            author='tester',title='',content='更新后闪退',rating=1,app_id='111',country='cn')
        values.update(changes)
        return RawReview(**values)

    def analyzed(self,reviews):
        save_reviews(self.db,reviews)
        return analyze_reviews(read_reviews(self.db,app_id='111',country='cn'),self.db)

    def test_negation_and_neutral_subjects(self):
        cases={'不好玩，不推荐':'差评','没有闪退，也不卡顿':'好评','服务器很稳定':'好评',
               'not good':'差评','no crashes':'待复核','沒有閃退，非常穩定':'好评','not bad':'好评'}
        for text,label in cases.items():
            with self.subTest(text=text): self.assertEqual(analyze_sentiment(text)['sentiment_label'],label)
        self.assertEqual(classify_issue('没有闪退，也不卡顿')['issue_category'],'未归类/待复核')
        self.assertEqual(classify_issue('日本的女性角色很可爱')['issue_category'],'内容/玩法反馈')

    def test_unsupported_and_conflicting_reviews_require_review(self):
        for text,rating in [('おすすめしません',1),('재미없어요',2),('闪退',5),('excellent',1),('👌',None)]:
            self.assertTrue(analyze_sentiment(text,rating)['needs_review'])
        self.assertEqual(analyze_sentiment('闪退',5)['sentiment_label'],'待复核')

    def test_same_text_without_identity_is_not_silently_dropped(self):
        self.assertEqual(len(prepare_dataframe(pd.DataFrame([{'content':'同样评论'},{'content':'同样评论'}])).data),2)

    def test_update_history_and_stale_override(self):
        data=self.analyzed([self.review()]); row=data.iloc[0]
        review_override(self.db,int(row.review_id),row.content_hash,'差评','Bug/闪退问题','QA','已核实')
        self.assertTrue(analyze_reviews(read_reviews(self.db),self.db).iloc[0].manual_reviewed)
        result=save_reviews(self.db,[self.review(content='修复后很好玩',rating=5,date='2026-09-22T02:00:00+00:00')])
        self.assertEqual(result.updated,1)
        current=analyze_reviews(read_reviews(self.db),self.db).iloc[0]
        self.assertTrue(current.override_stale)
        self.assertFalse(current.manual_reviewed)
        with database(self.db) as conn: self.assertEqual(conn.execute('SELECT COUNT(*) FROM review_revisions').fetchone()[0],1)
        with self.assertRaises(ValueError): review_override(self.db,int(row.review_id),row.content_hash,'差评','Bug/闪退问题','QA','旧页面')

    def test_old_revision_cannot_replace_new(self):
        self.analyzed([self.review(date='2026-09-22T03:00:00+00:00',rating=5)])
        save_reviews(self.db,[self.review(rating=1)])
        self.assertEqual(read_reviews(self.db).iloc[0].rating,5)

    def test_single_incident_is_not_overall_emergency(self):
        view=snapshot(self.analyzed([self.review()]),now=self.now)
        self.assertEqual(view['risk'],'样本不足')
        self.assertEqual(issue_plans(view)[0]['priority'],'P1')
        self.assertIn('未确认',issue_plans(view)[0]['trigger'])

    def test_healthy_and_empty_reports_do_not_invent_incidents(self):
        for data in [analyze_reviews(pd.DataFrame()),self.analyzed([self.review(content='很好玩，很稳定',rating=5)])]:
            view=snapshot(data,now=self.now)
            plans=issue_plans(view)
            self.assertEqual(plans,[])
            md,html,_=build_report({'app_id':'111','country':'cn','app_name':'健康测试'},view,plans,[])
            self.assertNotIn('需要系统性治理',md)
            self.assertNotIn('P0',md)

    def test_no_recent_data_does_not_relabel_old_data_as_recent(self):
        view=snapshot(self.analyzed([self.review(date='2026-06-01T00:00:00+00:00')]),now=self.now)
        self.assertEqual(view['metrics']['total'],0)

    def test_china_midnight_window_and_unknown_dates(self):
        rows=[self.review('a',date='2026-09-21T16:01:00+00:00'),self.review('b',date='2026-09-21T15:59:00+00:00'),self.review('c',date=None)]
        view=snapshot(self.analyzed(rows),days=1,now=self.now)
        self.assertEqual(view['metrics']['total'],1)
        self.assertEqual(view['undated'],1)

    def test_comparison_requires_samples_and_detects_rise(self):
        reviews=[self.review(str(i),rating=1 if i<20 else 5,content='一般',date='2026-09-22T01:00:00+00:00') for i in range(40)]
        reviews += [self.review('old'+str(i),rating=1 if i<2 else 5,content='一般',date='2026-09-14T01:00:00+00:00') for i in range(40)]
        view=snapshot(self.analyzed(reviews),now=self.now)
        self.assertEqual(view['delta'],45)
        self.assertEqual(view['risk'],'需要优先核查')

    def test_game_country_isolation(self):
        data=self.analyzed([self.review(),self.review(country='us'),self.review(app_id='222')])
        self.assertEqual(len(data),1)
        self.assertEqual(len(read_reviews(self.db)),3)

    def test_polling_lease_serializes_workers_and_retry(self):
        configure_target(self.db,'111','cn',True,15,1)
        request_collection(self.db,'111','cn')
        job=claim_job(self.db,'111','cn')
        self.assertIsNotNone(job)
        self.assertIsNone(claim_job(self.db,'111','cn'))
        collector=Mock(); collector.collect.return_value=CollectionResult(platform='App Store',errors=['网络超时'])
        self.assertEqual(perform_job(self.db,job,self.folder,collector),'失败')
        state=target_state(self.db,'111','cn')
        self.assertIsNone(state['last_success']); self.assertEqual(state['failures'],1)
        self.assertEqual(recent_runs(self.db,'111','cn')[0]['detail'],'网络超时')
        self.assertIsNone(claim_job(self.db,'111','cn'))
        request_collection(self.db,'111','cn')
        collector.collect.return_value=CollectionResult(platform='App Store',reviews=[self.review()],stop_reason='页数上限')
        perform_job(self.db,claim_job(self.db,'111','cn'),self.folder,collector)
        self.assertIsNotNone(target_state(self.db,'111','cn')['last_success'])

    def test_paused_monitor_does_not_run_without_request(self):
        configure_target(self.db,'111','cn',False)
        self.assertIsNone(claim_job(self.db,'111','cn'))
        request_collection(self.db,'111','cn')
        self.assertIsNotNone(claim_job(self.db,'111','cn'))

    def test_repeated_feed_page_stops_with_warning(self):
        payload=json.loads((Path(__file__).parent/'fixtures/appstore_reviews.json').read_text())
        response=Mock(status_code=200); response.json.return_value=payload
        session=Mock(); session.get.return_value=response
        collector=AppStoreCollector(self.folder,session=session)
        result=collector.collect('111',max_pages=10,delay_seconds=0,detect_repeated_pages=True)
        self.assertEqual(session.get.call_count,2)
        self.assertEqual(result.fetched_count,2)
        self.assertEqual(result.stop_reason,'重复页面')

    def test_import_checks_binding_and_preserves_id_precision(self):
        text='app_id,country,external_id,date,rating,content\n111,cn,123456789012345678,2026-09-01,1,闪退\n'
        rows,errors,_,_=parse_import(text.encode(),'test.csv','111','cn')
        self.assertFalse(errors); self.assertEqual(rows[0].external_id,'123456789012345678')
        self.assertTrue(rows[0].date.endswith('+08:00'))
        self.assertEqual(len(parse_import(text.encode(),'test.csv','222','cn')[1]),1)

    def test_import_rejects_invalid_values_and_preserves_anonymous_rows(self):
        raw='date,rating,content\n2026-09-01,1,闪退\n2026-09-01,1,闪退\n'
        rows,errors,_,_=parse_import(raw.encode(),'test.csv','111','cn')
        self.assertFalse(errors); self.assertNotEqual(rows[0].external_id,rows[1].external_id)
        self.assertEqual(save_reviews(self.db,rows).inserted,2)
        self.assertEqual(save_reviews(self.db,rows).duplicates,2)
        self.assertEqual(len(parse_import(raw.replace(',1,',',9,').encode(),'test.csv','111','cn')[1]),2)

    def test_excel_literal_and_html_escape(self):
        sheet=load_workbook(BytesIO(excel_bytes({'x':pd.DataFrame({'text':['=1+1','@x','-1+1']})})))['x']
        self.assertEqual(sheet['A2'].data_type,'s'); self.assertEqual(sheet['A2'].value,'=1+1')
        view=snapshot(self.analyzed([self.review(content='<script>alert(1)</script>闪退')]),now=self.now)
        _,html,_=build_report({'app_id':'111','country':'cn','app_name':'<b>test</b>'},view,issue_plans(view),[])
        self.assertNotIn('<script>',html); self.assertIn('&lt;script&gt;',html)

    def test_export_failure_preserves_old_and_unrelated_files(self):
        self.analyzed([self.review()])
        output=self.folder/'exports'; output.mkdir()
        untouched=output/'user_notes.xlsx'; untouched.write_bytes(b'keep')
        first=export_reviews_by_game(self.db,output)
        old=first[0].read_bytes()
        with patch('modules.safe_exports.excel_bytes',side_effect=RuntimeError('disk full')):
            with self.assertRaises(RuntimeError): export_reviews_by_game(self.db,output)
        self.assertEqual(untouched.read_bytes(),b'keep'); self.assertEqual(first[0].read_bytes(),old)

    def test_public_comment_control_characters_do_not_break_excel_or_change_source(self):
        source=pd.DataFrame({'评论':['第一行\n第二行\x14保留中文\t分隔','=1+1\x00']})
        sheet=load_workbook(BytesIO(excel_bytes({'评论':source})))['评论']
        self.assertEqual(sheet['A2'].value,'第一行\n第二行保留中文\t分隔')
        self.assertEqual(sheet['A3'].value,'=1+1')
        self.assertEqual(sheet['A3'].data_type,'s')
        self.assertIn('\x14',source.iloc[0,0])

    def test_action_lifecycle_dedup_and_audit(self):
        plan=issue_plans(snapshot(self.analyzed([self.review()]),now=self.now))[0]
        identity,created=create_action(self.db,'111','cn',plan)
        self.assertTrue(created); self.assertEqual(create_action(self.db,'111','cn',plan),(identity,False))
        action=actions_for(self.db,'111','cn')[0]
        with self.assertRaises(ValueError): update_action(self.db,identity,action['updated_at'],'已解决','','','', 'QA')
        update_action(self.db,identity,action['updated_at'],'观察效果','张三','2026-09-24','修复单 QA-123，24小时后回看','李四')
        with self.assertRaises(ValueError): update_action(self.db,identity,action['updated_at'],'已解决','张三','','完成','李四')
        with database(self.db) as connection: self.assertEqual(connection.execute('SELECT COUNT(*) FROM audit_log').fetchone()[0],2)

    def test_backup_is_complete_and_readable(self):
        self.analyzed([self.review()])
        path=backup_database(self.db)
        with database(path) as connection:
            self.assertEqual(connection.execute('PRAGMA quick_check').fetchone()[0],'ok')
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM review_records').fetchone()[0],1)

    def test_backup_restore_preserves_pre_restore_copy(self):
        from restore_backup import restore
        self.analyzed([self.review()])
        saved=backup_database(self.db)
        save_reviews(self.db,[self.review('second')])
        prior=restore(saved,self.db)
        self.assertEqual(len(read_reviews(self.db)),1)
        self.assertEqual(len(read_reviews(prior)),2)

    def test_alerts_deduplicate_and_acknowledge_without_closing_actions(self):
        from modules.alerts import publish_alerts,unread_alerts,acknowledge
        self.analyzed([self.review()])
        first=publish_alerts(self.db,'111','cn',now=self.now)
        self.assertGreater(first,0)
        self.assertEqual(publish_alerts(self.db,'111','cn',now=self.now),0)
        notices=unread_alerts(self.db)
        acknowledge(self.db,notices[0]['id'])
        self.assertEqual(len(unread_alerts(self.db)),first-1)
        self.assertEqual(actions_for(self.db,'111','cn'),[])

    def test_feed_fallback_keeps_storefront_and_parses_atom_next_page(self):
        payload=json.loads((Path(__file__).parent/'fixtures/appstore_reviews.json').read_text())
        payload['feed']['link']=[{'attributes':{'rel':'next','href':'https://itunes.apple.com/cn/rss/customerreviews/page=2/id=111/sortby=mostrecent/xml?l=en&urlDesc=x'}}]
        empty=Mock(status_code=200,text='{}'); empty.json.return_value={'feed':{}}
        first=Mock(status_code=200,text='{}'); first.json.return_value=payload
        second=Mock(status_code=200,text='''<?xml version="1.0"?><feed xmlns="http://www.w3.org/2005/Atom" xmlns:im="http://itunes.apple.com/rss"><entry><id>third</id><updated>2026-09-21T01:00:00Z</updated><title>Crash</title><content>Crashing again</content><im:rating>1</im:rating><im:version>3.0</im:version><author><name>Player</name></author></entry></feed>''')
        session=Mock(); session.get.side_effect=[empty,first,empty,second]
        result=AppStoreCollector(self.folder,session=session).collect('111','cn',max_pages=2,delay_seconds=0,follow_feed_links=True)
        self.assertEqual(result.fetched_count,3); self.assertFalse(result.errors)
        self.assertTrue(all('/cn/' in call.args[0] for call in session.get.call_args_list))
        self.assertEqual(result.reviews[-1].topic,'version:3.0')
        self.assertEqual(result.reviews[-1].external_id,'third')

    def test_feed_cannot_follow_different_game_or_region(self):
        payload=json.loads((Path(__file__).parent/'fixtures/appstore_reviews.json').read_text())
        payload['feed']['link']=[{'attributes':{'rel':'next','href':'https://itunes.apple.com/us/rss/customerreviews/page=2/id=222/json'}}]
        response=Mock(status_code=200,text='{}'); response.json.return_value=payload
        session=Mock(); session.get.return_value=response
        result=AppStoreCollector(self.folder,session=session).collect('111','cn',max_pages=2,delay_seconds=0,follow_feed_links=True)
        self.assertEqual(session.get.call_count,1); self.assertTrue(result.errors)

    def test_schema_upgrade_retains_reviews_and_creates_backup(self):
        self.analyzed([self.review()])
        with database(self.db) as connection:
            connection.execute("UPDATE store_meta SET value='2' WHERE key='schema_version'")
        init_db(self.db)
        self.assertEqual(len(read_reviews(self.db)),1)
        self.assertTrue(list((self.folder/'backups').glob('*.sqlite3')))

    def test_first_success_checks_existing_reviews_for_alerts(self):
        self.analyzed([self.review()])
        configure_target(self.db,'111','cn',False)
        request_collection(self.db,'111','cn')
        collector=Mock(); collector.collect.return_value=CollectionResult(platform='App Store',reviews=[self.review()])
        with patch('modules.alerts.publish_alerts') as publish:
            perform_job(self.db,claim_job(self.db,'111','cn'),self.folder,collector)
            publish.assert_called_once()

    def test_launcher_lock_rejects_duplicate_without_reading_locked_byte(self):
        from launch import acquire_lock
        path=self.folder/'instance.lock'
        first=acquire_lock(path)
        self.assertIsNotNone(first)
        try: self.assertIsNone(acquire_lock(path))
        finally: first.close()
        second=acquire_lock(path)
        self.assertIsNotNone(second)
        second.close()
